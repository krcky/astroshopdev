"""
Excel za proveru prevoda drugim modelom (Ivan, 2.10.2026) — iz JSON-a koji pravi `izvoz.ts`.

    npx tsx scripts/prevod/izvoz.ts > /tmp/prevod.json
    python3 scripts/prevod/excel.py /tmp/prevod.json Prevod-Astro-Shop.xlsx      (treba openpyxl)

List "Uputstvo" je tekst koji se kopira drugom modelu; po jedan list za svaki jezik. Model popunjava
kolone "ispravka" i "zasto"; `uvoz.py` vraca ispravke u recnik.
"""
import json
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: F401 (rezerva za buduce)

JEZICI = [
    ("hr", "Hrvatski", "hrvatski (Hrvatska): ijekavica, infinitiv umesto 'da + prezent', hrvatske reci (tko/što, tjedan, točka, račun, mobitel, postavke, spremi)"),
    ("bs", "Bosanski", "bosanski (BiH): ijekavica, recnik blizi srpskom (sedmica, hiljada, tačka, ko/šta, obavještenja), 'da + prezent' i infinitiv su oba u redu"),
    ("sl", "Slovenscina", "slovenacki: dvojina (1 dan, 2 dneva, 3 dni, 5 dni), Mesec = Luna, ascendent, ura rojstva, znamenje"),
    ("mk", "Makedonski", "makedonski: cirilica, clan po rodu (твојот/твојата/твоето), Месечина; 'AI' i 'Premium' namerno ostaju latinicom"),
    ("en", "English", "engleski (americki pravopis): sentence case, 'you', astroloski izrazi kako ih pisu engleski astrolozi"),
]

UPUTSTVO = """ZADATAK: Ti si izvorni govornik i lektor za jezik ovog lista. Proveravas prevod korisnickog interfejsa
mobilne aplikacije "Astro Shop" (horoskop). Prevod je napravio drugi AI i vec ga je jednom pregledao treci —
trazis ono sto su oba propustila. NE prevodis iznova; diras samo ono sto je pogresno ili ocigledno neprirodno.

KOLONE:
- kljuc: tehnicko ime (ne menjaj).
- gde se vidi: ekran ili komponenta i napomena (dugme, naslov, verzal, citac ekrana, kartica za deljenje...).
- srpski: ORIGINAL — on je istina za ZNACENJE.
- prevod: ono sto korisnik sada vidi.
- vazno: DA = prvi ekrani, pretplata, pristanak, greske — tu greska najvise kosta.
- ispravka: UPISI SAMO kad nesto treba promeniti — ceo nov tekst stavke. Prazno = u redu je.
- zasto: jedna kratka recenica.

STA TRAZIS (redom vaznosti):
1. Promenjeno ili izgubljeno znacenje; obecanje koje original ne daje (pretplata, placanje, pristanak, greske).
2. Gramatika: padez, rod, slaganje, mnozina uz broj, glagolski oblik, red reci.
3. Rec koja nije iz tog jezika ili isti pojam napisan na dva nacina na raznim mestima.
4. Neprirodno, "prevedeno" zvuci; predugo za dugme (~20 znakova) ili naslov (~30).
5. Pravopis i interpunkcija.

PRAVILA:
- Ton: na "ti", aplikacija TUMACI a ne prorice, bez uzvicnika, bez laskanja, kratke recenice.
- Ne prevodi: "Astro Shop", "astroshop.rs", imena (Boban Vujović), App Store, Google Play.
- Redovi koji pocinju sa "ƒ" su RECENICE SA PROMENLJIVOM. Deo `${...}` je mesto gde aplikacija ubacuje
  vrednost (ime, broj, datum, ime znaka). U ispravci ZADRZI svaki `${...}` tacno kako je napisan — samo
  promeni reci oko njega. Ako vise sablona stoji spojeno sa "…", to su varijante iste stavke (npr. za
  jedninu i mnozinu); ispravku napisi za svaku, istim redom, razdvojene sa "…".
- Redovi sa " | " su spiskovi (npr. imena meseci): ispravka je ceo spisak, istim redom, sa " | ".
- Odgovor astrologa je glasovna poruka NA SRPSKOM: u engleskom, slovenackom i makedonskom to mora da pise;
  u hrvatskom i bosanskom ne treba.

VRATI: isti fajl, sa popunjenim kolonama "ispravka" i "zasto" samo tamo gde nesto menjas.
"""

KOLONE = ["kljuc", "gde se vidi", "srpski", "prevod", "vazno", "ispravka", "zasto"]
SIRINE = [34, 46, 52, 52, 8, 52, 36]


def main(ulaz: str, izlaz: str) -> None:
    redovi = json.load(open(ulaz, encoding="utf-8"))
    wb = Workbook()
    u = wb.active
    u.title = "Uputstvo"
    u.column_dimensions["A"].width = 120
    for i, red in enumerate(UPUTSTVO.strip().split("\n"), start=1):
        c = u.cell(row=i, column=1, value=red)
        c.alignment = Alignment(wrap_text=True)
        if red.isupper() or red.endswith(":") and red.split(":")[0].isupper():
            c.font = Font(bold=True)
    # posebna napomena po jeziku
    start = len(UPUTSTVO.strip().split("\n")) + 2
    u.cell(row=start, column=1, value="JEZICI (po jedan list):").font = Font(bold=True)
    for i, (_, ime, opis) in enumerate(JEZICI, start=1):
        u.cell(row=start + i, column=1, value=f"- {ime}: {opis}")

    zaglavlje = PatternFill("solid", fgColor="E8E6F8")
    vazno = PatternFill("solid", fgColor="FFF4D6")
    for kod, ime, opis in JEZICI:
        ws = wb.create_sheet(ime)
        ws.append(KOLONE)
        for c in ws[1]:
            c.font = Font(bold=True)
            c.fill = zaglavlje
        # vazni prvi, pa ostali — unutar grupe redosled iz recnika
        for r in sorted(redovi, key=lambda r: not r["vazno"]):
            ws.append([r["kljuc"], r["kontekst"], r["sr"], r[kod], "DA" if r["vazno"] else "", "", ""])
            if r["vazno"]:
                for c in ws[ws.max_row][:5]:
                    c.fill = vazno
        for i, s in enumerate(SIRINE, start=1):
            ws.column_dimensions[chr(64 + i)].width = s
        for row in ws.iter_rows(min_row=2):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:G{ws.max_row}"
        ws.sheet_properties.tabColor = "6E6AE8"
        print(f"{ime}: {ws.max_row - 1} redova ({sum(r['vazno'] for r in redovi)} vaznih) — {opis[:40]}…")
    wb.save(izlaz)
    print(f"-> {izlaz}")


if __name__ == "__main__":
    main(sys.argv[1], sys.argv[2])
